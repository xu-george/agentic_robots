from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from PIL import Image

OUT = Path(__file__).resolve().parent
PPTX = OUT / "unified_robot_control_interface.pptx"

# 16:9 canvas
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

# Light-blue presentation theme.
BG = RGBColor(239, 247, 255)
PANEL = RGBColor(255, 255, 255)
PANEL2 = RGBColor(220, 237, 252)
TEXT = RGBColor(19, 43, 68)
MUTED = RGBColor(79, 105, 130)
CYAN = RGBColor(0, 137, 145)
BLUE = RGBColor(35, 105, 190)
GOLD = RGBColor(181, 119, 0)
GREEN = RGBColor(25, 135, 79)
RED = RGBColor(196, 55, 71)
PURPLE = RGBColor(112, 67, 175)
WHITE = RGBColor(255, 255, 255)
DARK = RGBColor(226, 240, 253)

W, H = 13.333, 7.5


def rect(slide, x, y, w, h, fill=PANEL, line=None, radius=True, transparency=0):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = fill; shp.fill.transparency = transparency
    shp.line.color.rgb = line or fill
    return shp


def txt(slide, text, x, y, w, h, size=18, color=TEXT, bold=False,
        font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.04):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.name = font; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return box


def rich(slide, runs, x, y, w, h, size=18, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True; tf.vertical_anchor = valign
    tf.margin_left = tf.margin_right = Inches(.04); tf.margin_top = tf.margin_bottom = Inches(.02)
    p = tf.paragraphs[0]; p.alignment = align
    for s, c, b in runs:
        r = p.add_run(); r.text = s; r.font.name = "Aptos"; r.font.size = Pt(size); r.font.color.rgb = c; r.font.bold = b
    return box


def line(slide, x1, y1, x2, y2, color=MUTED, width=1.5, dash=None, arrow=False):
    s = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    s.line.color.rgb = color; s.line.width = Pt(width)
    if dash: s.line.dash_style = dash
    # python-pptx does not currently expose connector arrowheads portably.
    # Direction is conveyed by left-to-right layout and terminal state labels.
    return s


def chip(slide, label, x, y, w, color=CYAN):
    rect(slide, x, y, w, .3, fill=color, line=color, transparency=82)
    txt(slide, label, x+.08, y+.03, w-.16, .22, 10, color, True, align=PP_ALIGN.CENTER)


def base(title, section, source=""):
    s = prs.slides.add_slide(blank)
    bg = s.background.fill; bg.solid(); bg.fore_color.rgb = BG
    txt(s, section.upper(), .55, .28, 3.2, .25, 10, CYAN, True)
    txt(s, title, .55, .62, 12.1, .52, 27, TEXT, True)
    line(s, .55, 1.23, 12.78, 1.23, RGBColor(45, 58, 75), 1)
    n = len(prs.slides)
    txt(s, f"{n:02d}", 12.55, 7.12, .25, .18, 9, MUTED, True, align=PP_ALIGN.RIGHT)
    if source:
        txt(s, source, .55, 7.10, 11.6, .22, 8, MUTED)
    return s


def card(slide, title, body, x, y, w, h, accent=CYAN, tag=None):
    rect(slide, x, y, w, h, PANEL, RGBColor(46, 60, 78))
    rect(slide, x, y, .06, h, accent, accent, radius=False)
    txt(slide, title, x+.22, y+.18, w-.42, .34, 16, TEXT, True)
    txt(slide, body, x+.22, y+.67, w-.42, h-.82, 12.5, MUTED)
    if tag: chip(slide, tag, x+w-1.25, y+.17, 1.02, accent)


def add_bullets(slide, items, x, y, w, h, size=16, color=TEXT, bullet_color=CYAN, gap=8):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.02); tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap); p.level = 0
        r = p.add_run(); r.text = "●  "; r.font.color.rgb = bullet_color; r.font.size = Pt(size-2)
        r = p.add_run(); r.text = item; r.font.name = "Aptos"; r.font.size = Pt(size); r.font.color.rgb = color
    return box


def picture_crop(slide, path, x, y, w, h):
    path = str(path)
    im = Image.open(path); iw, ih = im.size
    target = w / h; source = iw / ih
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    if source > target:
        crop = (1 - target/source)/2; pic.crop_left = pic.crop_right = crop
    elif source < target:
        crop = (1 - source/target)/2; pic.crop_top = pic.crop_bottom = crop
    return pic

# 1 — title
s = prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
rect(s, 8.72, 0, 4.61, 7.5, PANEL, PANEL, radius=False)
for i, (c, yy, ww) in enumerate([(CYAN,1.05,3.8),(BLUE,1.52,3.05),(GOLD,1.99,2.3)]):
    rect(s, 9.12, yy, ww, .12, c, c, radius=False)
txt(s, "UNIFIED ROBOT\nCONTROL INTERFACE", .65, 1.25, 7.7, 1.45, 36, TEXT, True)
txt(s, "A common GUI and integration framework for agentic planning, Isaac ROS execution, and Picker2 learned skills", .68, 3.02, 7.2, 1.0, 19, MUTED)
chip(s, "AGENTIC", .68, 4.42, 1.35, CYAN); chip(s, "ROS 2", 2.18, 4.42, 1.1, BLUE); chip(s, "RL SKILLS", 3.43, 4.42, 1.35, PURPLE)
txt(s, "Architecture proposal • grounded in the current repositories", .68, 6.75, 7.5, .28, 11, MUTED)
# abstract robot stack glyph
for yy, label, c in [(3.0,"PLAN",CYAN),(3.75,"PERCEIVE",BLUE),(4.5,"ACT",GOLD),(5.25,"VERIFY",GREEN)]:
    rect(s, 9.35, yy, 3.2, .55, PANEL2, c)
    txt(s, label, 9.55, yy+.14, 2.8, .22, 13, c, True, align=PP_ALIGN.CENTER)

# 2 — full CaP-X simulation screenshot supplied by the user
s = base("CaP-X simulation: workflow and 3D visualization", "Interface example", "Screenshot captured from the local Franka cube-lifting demonstration")
latest_shot = OUT / "assets" / "capx_latest_screenshot.png"
if latest_shot.exists():
    picture_crop(s, latest_shot, .62, 1.42, 12.08, 5.42)
    rect(s, .82, 6.42, 4.55, .28, WHITE, WHITE)
    txt(s, "Executable workflow • task events • MuJoCo view • interactive Viser scene", .92, 6.48, 4.3, .15, 8.5, TEXT, True)

# 3 — explain the three principal GUI surfaces
s = base("How the CaP-X GUI exposes robot execution", "GUI structure", "Reference: local CaP-X Franka cube-lifting demonstration")
card(s, "1  Pseudo-code interface", "High-level, readable robot workflow.\n\n• Object pose / grasp pose\n• Motion planning\n• Robot and gripper commands\n• Execution result", .55, 1.5, 3.95, 4.65, CYAN, "WORKFLOW")
card(s, "2  YAML task configuration", "Reproducible task setup.\n\n• Robot and simulator\n• APIs and task prompt\n• Model / workflow\n• Visualization and outputs", 4.69, 1.5, 3.95, 4.65, BLUE, "CONFIG")
card(s, "3  3D visualization", "Live execution context.\n\n• Robot and object state\n• Camera view and frames\n• Grasps and trajectories\n• Interactive inspection", 8.83, 1.5, 3.95, 4.65, PURPLE, "VISER")
rect(s, 1.05, 6.42, 11.25, .46, DARK, RGBColor(88, 125, 160))
rich(s, [("Task YAML", BLUE, True), ("  →  high-level Python workflow", CYAN, True), ("  →  perception / planning / motion services", GOLD, True), ("  →  live 3D evidence", PURPLE, True)], 1.25, 6.54, 10.85, .2, 11.5, PP_ALIGN.CENTER)

# 4 — API modules and the implementations that currently provide them
s = base("Agentic robot APIs and current implementations", "Module view", "Sources: agentic_robots/apis · agentic_robots/services · Isaac ROS manipulation workflow")
# Top-level execution flow
for i,(label,c) in enumerate([("Task intent",CYAN),("CodeAgent",PURPLE),("Allowlisted runtime",GOLD),("Robot capabilities",BLUE)]):
    x=.62+i*3.05
    rect(s,x,1.48,2.62,.54,PANEL2,c)
    txt(s,label,x+.1,1.63,2.42,.22,13,c,True,align=PP_ALIGN.CENTER)
    if i<3: line(s,x+2.62,1.75,x+3.05,1.75,MUTED,2)
# Three API cards
card(s, "Robot Control API", "Interface\n• pose and robot state\n• Cartesian/joint motion\n• gripper open/close\n\nCurrent implementation\n• ActionApi\n• FrankaReal + libfranka bridge\n• FrankaSim / PyBullet", .62, 2.35, 3.72, 3.78, CYAN, "ACTION")
card(s, "Perception API", "Interface\n• object segmentation and pose\n• point-cloud extraction\n• grasp candidates\n\nCurrent implementation\n• SAM3 segmentation service\n• Contact-GraspNet grasp service\n• calibrated D455 observations", 4.8, 2.35, 3.72, 3.78, BLUE, "PERCEPTION")
card(s, "Planning API", "Interface\n• IK and target-pose planning\n• collision-aware trajectory\n• execute / cancel / report\n\nCurrent implementation\n• PyRoKi motion service (experimental)\n• cuMotion + nvblox available in Isaac ROS\n• unified PlanningApi adapter: proposed", 8.98, 2.35, 3.72, 3.78, PURPLE, "PLANNING")
rect(s,.95,6.43,11.45,.42,DARK,RGBColor(88,125,160))
rich(s,[("Current path: ",TEXT,True),("SAM3 → mask",BLUE,True),("  •  GraspNet → grasp",CYAN,True),("  •  PyRoKi / cuMotion → trajectory",PURPLE,True),("  •  robot backend → motion",GOLD,True)],1.15,6.54,11.05,.2,10.5,PP_ALIGN.CENTER)

# 5
s = base("Preserve small contracts at the planner boundary", "Interfaces", "Sources: agentic_robots/apis/action/api.py · apis/perception/api.py · agent/runtime.py")
card(s, "ActionApi", "get_end_effector_pose()\nmove_to_pose(position, quaternion_wxyz)\nopen_gripper() / close_gripper()", .62, 1.55, 3.72, 2.15, CYAN)
card(s, "PerceptionApi", "get_observation() • segment_object()\nget_object_pose() • plan_grasp()\nget_finger_grasp()", 4.8, 1.55, 3.72, 2.15, BLUE)
card(s, "Runtime envelope", "Only registered functions and a small builtin set. No arbitrary imports or system access from generated task code.", 8.98, 1.55, 3.72, 2.15, GOLD)
rect(s, .62, 4.15, 12.08, 1.75, DARK, RGBColor(50,66,84))
txt(s, "Recommended extension", .9, 4.4, 2.35, .28, 13, GOLD, True)
add_bullets(s, ["Wrap ROS actions and RL policies behind the same capability registry.", "Keep frames, units, preconditions, timeout, cancellation, and evidence explicit."], 3.0, 4.38, 9.2, 1.15, 15)
chip(s, "IMPLEMENTED", .7, 6.3, 1.35, GREEN); txt(s, "current repository behavior", 2.15, 6.34, 2.2, .2, 10, MUTED)
chip(s, "PROPOSED", 5.05, 6.3, 1.2, GOLD); txt(s, "integration contract", 6.35, 6.34, 2.2, .2, 10, MUTED)

# 6
s = base("Isaac ROS supplies the production ROS graph", "Isaac ROS", "Source: isaac_ros_manipulation_bringup/launch/workflows.launch.py and interfaces/action/*.action")
# flow
nodes=[("RGB-D / sim",CYAN),("Drivers + TF",BLUE),("nvblox",GREEN),("Detection / pose",PURPLE),("cuMotion",GOLD),("Pick & place",RED)]
for i,(lab,c) in enumerate(nodes):
    x=.55+i*2.12; rect(s,x,1.62,1.75,.68,PANEL2,c); txt(s,lab,x+.08,1.82,1.59,.24,12,c,True,align=PP_ALIGN.CENTER)
    if i<5: line(s,x+1.75,1.96,x+2.12,1.96,MUTED,2)
card(s, "Available orchestration surfaces", "ROS 2 launch configurations; manipulation servers; actions for segmentation, detection, pose, PickAndPlace, and MultiObjectPickAndPlace.", .62, 2.78, 5.76, 2.1, BLUE)
card(s, "Integration principle", "The unified gateway should be a ROS 2 client—not a replacement for ROS lifecycle, TF, actions, drivers, or real-time controllers.", 6.72, 2.78, 5.98, 2.1, GOLD)
rect(s, .62, 5.35, 12.08, .78, PANEL, RGBColor(47,61,79))
rich(s, [("GUI asks: ",CYAN,True),("“execute capability + show evidence”",TEXT,False),("   ROS decides: ",BLUE,True),("planning, feedback, cancellation, result.",TEXT,False)], .9, 5.6, 11.5, .3, 15, PP_ALIGN.CENTER)

# 7
s = base("Borrow CaP-X’s interaction pattern—not its coupling", "GUI reference", "Sources: cap-x/web-ui/src/App.tsx · capx/web/server.py · capx/web/models.py")
# Screenshot captured from the locally running CaP-X Vite/FastAPI application.
capx_shot = OUT / "assets" / "capx_cube_lift_3d_live.png"
if capx_shot.exists():
    picture_crop(s,capx_shot,.55,1.45,8.0,4.95)
    rect(s,.55,6.04,3.18,.28,DARK,DARK)
    txt(s,"Actual local CaP-X UI • idle state",.68,6.1,2.9,.15,8,TEXT,True)
card(s,"Patterns worth keeping","• Config + model selection\n• Start / stop / reset\n• Session reconnect and event replay\n• Streaming reasoning + code + results\n• Human feedback between turns\n• Embedded 3D view",8.85,1.45,3.85,2.8,CYAN)
card(s,"Changes required for robots","Add hardware state, safety mode, approvals, alarms, camera/TF overlays, capability health, and an immutable audit timeline.",8.85,4.52,3.85,1.88,GOLD)

# 8
s = base("Proposed unified operator workspace", "GUI concept", "Concept mock-up informed by the CaP-X web UI")
# full mock UI
rect(s,.42,1.38,12.48,5.55,DARK,RGBColor(53,69,89))
rect(s,.42,1.38,12.48,.56,PANEL2,PANEL2,radius=False)
txt(s,"ROBOT OPS",.65,1.56,1.5,.2,12,TEXT,True); chip(s,"SIM",2.15,1.51,.72,PURPLE); chip(s,"ARMED",10.47,1.51,.9,GREEN); chip(s,"E-STOP",11.5,1.51,1.02,RED)
# left nav
rect(s,.42,1.94,1.1,4.99,PANEL,PANEL,radius=False)
for i,(lab,c) in enumerate([("TASK",CYAN),("SCENE",BLUE),("SKILLS",PURPLE),("LOGS",GOLD),("HEALTH",GREEN)]):
    txt(s,lab,.56,2.25+i*.67,.82,.2,9,c,True,align=PP_ALIGN.CENTER)
# center timeline/chat
rect(s,1.52,1.94,4.47,4.99,BG,BG,radius=False)
txt(s,"TASK  Load suitcase 101050",1.78,2.16,3.9,.3,13,CYAN,True)
for yy,label,c in [(2.7,"01  Detect handle",GREEN),(3.25,"02  Plan approach",GREEN),(3.8,"03  Execute RL grasp",GOLD),(4.35,"04  Transfer to platform",MUTED),(4.9,"05  Verify stable release",MUTED)]:
    rect(s,1.78,yy,3.93,.4,PANEL2,RGBColor(46,61,79)); txt(s,label,1.94,yy+.1,3.55,.18,11,c,True)
rect(s,1.78,6.3,3.93,.34,PANEL2,PANEL2); txt(s,"Operator feedback / approval…",1.92,6.37,3.5,.18,9,MUTED)
# viz
rect(s,5.99,1.94,4.6,3.5,PANEL,PANEL,radius=False); txt(s,"3D + RGB-D + overlays",6.25,2.17,3.2,.25,12,BLUE,True)
# robot geometric
line(s,7.15,4.65,7.65,3.8,GOLD,7); line(s,7.65,3.8,8.45,3.25,GOLD,7); line(s,8.45,3.25,9.15,3.63,GOLD,7)
rect(s,8.95,3.55,.55,.2,CYAN,CYAN); rect(s,6.65,4.65,1.2,.25,PANEL2,PANEL2)
# evidence/telemetry
rect(s,5.99,5.44,4.6,1.49,PANEL2,PANEL2,radius=False); txt(s,"Evidence",6.25,5.67,1.2,.23,11,GREEN,True)
txt(s,"mask 0.91   grasp 0.74   TF ✓\nphase GRASP   policy v4   42 ms",6.25,6.02,3.9,.55,10,MUTED,font="Aptos Mono")
# right rail
rect(s,10.59,1.94,2.31,4.99,PANEL,PANEL,radius=False); txt(s,"SYSTEM",10.85,2.17,1.75,.24,11,TEXT,True)
for i,(lab,val,c) in enumerate([("Robot","ready",GREEN),("ROS graph","healthy",GREEN),("Perception","3/3",GREEN),("Planner","online",GREEN),("Policy","loaded",PURPLE),("Safety","guarded",GOLD)]):
    txt(s,lab,10.83,2.7+i*.55,1.0,.18,9,MUTED); txt(s,val,11.78,2.7+i*.55,.8,.18,9,c,True,align=PP_ALIGN.RIGHT)

# 9
s = base("Reference architecture: control plane + adapters", "Target architecture")
# columns
for x,title,c in [(.52,"WEB CLIENT",CYAN),(3.05,"CONTROL GATEWAY",GOLD),(6.12,"CAPABILITY ADAPTERS",BLUE),(10.12,"EXECUTION",PURPLE)]:
    txt(s,title,x,1.45,2.3,.25,11,c,True,align=PP_ALIGN.CENTER)
rect(s,.62,1.9,2.1,3.85,PANEL,RGBColor(47,62,81));
for i,t in enumerate(["Task builder","Live timeline","Scene / video","Safety controls","Audit viewer"]): txt(s,t,.86,2.25+i*.63,1.62,.23,12,TEXT,True,align=PP_ALIGN.CENTER)
rect(s,3.18,1.9,2.55,3.85,PANEL,RGBColor(47,62,81));
for i,(t,c) in enumerate([("Session manager",GOLD),("Capability registry",GOLD),("Event broker",CYAN),("Policy/approval",RED),("Artifact store",GREEN)]): rect(s,3.47,2.17+i*.63,1.97,.38,PANEL2,c); txt(s,t,3.57,2.27+i*.63,1.77,.18,10,c,True,align=PP_ALIGN.CENTER)
rect(s,6.23,1.9,3.5,3.85,PANEL,RGBColor(47,62,81));
for i,(t,c) in enumerate([("Agentic Python API",CYAN),("ROS 2 action/service",BLUE),("RL policy skill",PURPLE),("Simulation env",GREEN),("Camera / visualization",GOLD)]): rect(s,6.55,2.17+i*.63,2.85,.38,PANEL2,c); txt(s,t,6.67,2.27+i*.63,2.6,.18,10,c,True,align=PP_ALIGN.CENTER)
rect(s,10.24,1.9,2.45,3.85,PANEL,RGBColor(47,62,81));
for i,t in enumerate(["Isaac ROS","Franka bridge","Genesis","PyBullet","Model services"]): txt(s,t,10.55,2.25+i*.63,1.83,.23,12,TEXT,True,align=PP_ALIGN.CENTER)
for x1,x2 in [(2.72,3.18),(5.73,6.23),(9.73,10.24)]: line(s,x1,3.82,x2,3.82,CYAN,2,arrow=True)
rect(s,2.35,6.15,8.65,.58,DARK,RGBColor(48,65,84)); txt(s,"WebSocket event stream  •  REST commands  •  ROS 2 actions  •  typed artifacts",2.65,6.33,8.05,.22,13,MUTED,True,align=PP_ALIGN.CENTER)

# 10
s = base("Normalize every backend into a capability contract", "Control contract")
# pseudo-schema
rect(s,.62,1.48,5.2,4.95,DARK,RGBColor(52,68,88))
txt(s,"Capability",.92,1.75,1.2,.25,13,CYAN,True)
txt(s,"id: picker2.load_luggage\nmode: policy | action | agent\ninputs: object_id, target_pose\nframes: base_link → grasp_frame\npreconditions: robot_ready, tf_valid\ntimeout_s: 120\ncancel: required\nsafety_class: supervised\nevidence: rgb, depth, trajectory, result",.92,2.18,4.55,3.8,14,TEXT,font="Aptos Mono")
card(s,"Command envelope","task_id • session_id • capability_id • typed inputs • frame IDs • units • deadline • approval token",6.18,1.48,6.5,1.42,BLUE)
card(s,"Progress events","QUEUED → PLANNING → WAITING_APPROVAL → EXECUTING → VERIFYING → SUCCEEDED / FAILED / CANCELLED",6.18,3.12,6.5,1.42,GOLD)
card(s,"Result + evidence","Success code, failure reason, metrics, output poses, logs, images, trajectory/policy versions, and timestamps.",6.18,4.76,6.5,1.42,GREEN)

# 11
s = base("Make task execution observable and interruptible", "Lifecycle")
states=[("DRAFT",MUTED),("VALIDATE",CYAN),("PLAN",BLUE),("APPROVE",GOLD),("EXECUTE",PURPLE),("VERIFY",GREEN),("DONE",GREEN)]
for i,(lab,c) in enumerate(states):
    x=.46+i*1.84; rect(s,x,2.0,1.43,.62,PANEL2,c); txt(s,lab,x+.05,2.19,1.33,.2,11,c,True,align=PP_ALIGN.CENTER)
    if i<6: line(s,x+1.43,2.31,x+1.84,2.31,MUTED,2,arrow=True)
line(s,9.38,2.62,9.38,3.35,RED,2,arrow=True)
rect(s,8.55,3.35,1.68,.58,PANEL2,RED); txt(s,"STOP / SAFE",8.67,3.53,1.45,.2,11,RED,True,align=PP_ALIGN.CENTER)
card(s,"Verification is explicit","Task success is not “code ran.” Use environment success callbacks, ROS action results, contact/pose dwell checks, and operator confirmation where required.",.62,4.4,5.75,1.55,GREEN)
card(s,"Cancellation is end-to-end","UI stop → gateway cancellation → ROS action cancel / policy halt → controller stop → safe state → incident event.",6.72,4.4,5.96,1.55,RED)

# 12
s = base("Picker2 adds a learned, contact-rich luggage skill", "Picker2 task", "Sources: picker2_rl_grasping/README.md · envs/picker2.py · envs/luggage_env.py")
img=Path("/home/user/Documents/Code/picker2_rl_grasping/artifacts/parallel_rl_8x8.png")
if img.exists():
    picture_crop(s,img,.62,1.5,5.25,4.72)
    rect(s,.78,5.68,2.55,.34,DARK,DARK); txt(s,"64 parallel Genesis environments",.9,5.76,2.32,.18,9,TEXT,True)
card(s,"Current implementation","• Fixed Picker2 base, Doosan M1013, 2FG14\n• Physical hook-fingertip grasp; no attachment constraint\n• 7D delta pose + gripper action\n• 46D privileged observation\n• Parallel independent reset/step",6.2,1.5,6.45,2.2,PURPLE,"BUILT")
card(s,"Task sequence","REACH → GRASP → LIFT → TRANSPORT → ROTATE FLAT → PLACE → RELEASE → RETREAT → READY",6.2,3.95,6.45,1.25,CYAN)
rect(s,6.2,5.45,6.45,.75,DARK,RGBColor(52,68,88)); rich(s,[("Important: ",GOLD,True),("training and real-robot deployment remain incomplete; expose status honestly in the GUI.",TEXT,False)],6.48,5.65,5.9,.34,13)

# 13
s = base("Integrate Picker2 as a versioned skill adapter", "Picker2 integration")
# flow rows
rows=[("1  Perceive", "D455 / Isaac ROS pose or agentic segmentation", BLUE),
      ("2  Select", "policy version + suitcase model + target", CYAN),
      ("3  Guard", "workspace, TF, readiness, confidence, approval", GOLD),
      ("4  Execute", "policy runner publishes bounded 7D commands", PURPLE),
      ("5  Monitor", "phase, contacts, pose error, timeout, anomalies", GREEN),
      ("6  Verify", "flat + in bounds + stable + released + ready", RED)]
for i,(a,b,c) in enumerate(rows):
    y=1.48+i*.82; rect(s,.62,y,2.05,.58,PANEL2,c); txt(s,a,.78,y+.17,1.72,.2,12,c,True)
    line(s,2.67,y+.29,3.05,y+.29,c,2,arrow=True)
    rect(s,3.05,y,5.35,.58,PANEL,RGBColor(46,61,79)); txt(s,b,3.3,y+.16,4.86,.22,12,TEXT)
    rect(s,8.68,y,3.98,.58,DARK,RGBColor(46,61,79));
    event=["pose + confidence","skill manifest","approval event","action telemetry","phase evidence","task result"][i]
    txt(s,event,8.9,y+.17,3.55,.2,11,MUTED,True,align=PP_ALIGN.CENTER)
rect(s,.62,6.5,12.04,.48,DARK,GOLD)
rich(s,[("Deployment boundary: ",GOLD,True),("deterministic inference runner with manifest, normalization, action scaling, watchdog, reset, and stop hooks—never the training environment inside the GUI server.",TEXT,False)],.85,6.62,11.58,.24,10.5,PP_ALIGN.CENTER)

# 14
s = base("Safety must be independent of the agent and GUI", "Safety model")
# concentric-ish layers
rect(s,.7,1.5,11.95,4.95,DARK,RED)
rect(s,1.15,1.92,11.05,4.08,PANEL,GOLD)
rect(s,1.62,2.34,10.11,3.2,PANEL2,BLUE)
rect(s,2.12,2.78,9.1,2.3,BG,CYAN)
txt(s,"OPERATOR UI",.9,1.62,2.2,.25,11,RED,True)
txt(s,"Gateway authorization + approvals + audit",1.42,2.05,4.2,.25,11,GOLD,True)
txt(s,"Capability checks: schemas • TF • limits • timeouts",1.9,2.47,4.7,.25,11,BLUE,True)
txt(s,"ROS controller / hardware interlocks / user stop",2.42,2.93,5.1,.25,12,CYAN,True)
add_bullets(s,["The browser emergency control is a request, never the sole safety channel.","Generated code and learned policies only issue bounded high-level commands.","Real mode requires positive readiness, explicit arming, and clear mode indication."],2.5,3.45,8.5,1.28,14,TEXT,RED,5)
chip(s,"NON-NEGOTIABLE",5.45,5.66,2.3,RED)

# 15
s = base("Deliver in four increments", "Roadmap")
phases=[("01","OBSERVE","Unify camera, 3D, health, logs; read-only sessions.",CYAN),
        ("02","COMMAND","Capability registry, task lifecycle, stop/reset, ROS actions.",BLUE),
        ("03","AGENT","CodeAgent trace, approvals, retries, evidence and replay.",GOLD),
        ("04","PICKER2","Policy runner, sim evaluation, shadow mode, supervised hardware.",PURPLE)]
for i,(num,name,body,c) in enumerate(phases):
    x=.55+i*3.16
    rect(s,x,1.58,2.87,4.72,PANEL,RGBColor(48,63,82))
    txt(s,num,x+.2,1.84,.55,.4,24,c,True)
    txt(s,name,x+.2,2.42,2.4,.28,15,TEXT,True)
    txt(s,body,x+.2,2.92,2.43,1.55,12.5,MUTED)
    milestones=["contract","events","adapter","gate"][i]
    chip(s,milestones.upper(),x+.45,5.45,1.97,c)
# acceptance strip
rect(s,.55,6.55,12.08,.38,DARK,RGBColor(50,67,87)); txt(s,"Gate each phase on replayable evidence, cancellation tests, and simulator-first acceptance.",.85,6.64,11.5,.18,11,GREEN,True,align=PP_ALIGN.CENTER)

# 16
s = base("Recommended decision", "Conclusion")
txt(s,"Build the GUI as a thin, event-driven control plane over typed capabilities.",.8,1.55,11.75,.72,26,TEXT,True,align=PP_ALIGN.CENTER)
card(s,"Keep","agentic_robots API/runtime boundaries; Isaac ROS action semantics; Picker2’s single batched action/state interface; CaP-X session + split-view interaction model.",.7,2.62,3.75,2.55,GREEN)
card(s,"Add","capability registry; normalized lifecycle/events; ROS and policy adapters; evidence store; approval and safety policy; hardware status rail.",4.78,2.62,3.75,2.55,CYAN)
card(s,"Avoid","GUI-owned robot logic; direct browser-to-controller commands; training-code imports in deployment; ambiguous frames/units; “success = process exited.”",8.86,2.62,3.75,2.55,RED)
rect(s,2.18,5.72,8.98,.68,DARK,GOLD); txt(s,"First concrete deliverable: read-only unified session viewer + capability schema.",2.4,5.94,8.54,.24,15,GOLD,True,align=PP_ALIGN.CENTER)

# 17 sources
s = base("Repository evidence used in this proposal", "Appendix")
cols=[
("agentic_robots",["README.md","docs/project/STRUCTURE.md","agent/code_agent.py","agent/runtime.py","apis/action/api.py","apis/perception/api.py","web/realsense_viewer.py"]),
("Isaac ROS",["isaac_ros_manipulation/README.md","bringup/launch/workflows.launch.py","interfaces/action/PickAndPlace.action","interfaces/action/MultiObjectPickAndPlace.action","servers/launch/*.launch.py"]),
("Picker2 + CaP-X",["picker2_rl_grasping/README.md","plan.md","envs/picker2.py","envs/luggage_env.py","cap-x/web-ui/src/App.tsx","capx/web/server.py","capx/web/models.py"])]
for i,(head,items) in enumerate(cols):
    x=.55+i*4.22; card(s,head,"\n".join(items),x,1.5,3.94,4.95,[CYAN,BLUE,PURPLE][i])
txt(s,"Labels in this deck distinguish current implementation from proposed integration where relevant.",.65,6.67,12,.26,11,MUTED,align=PP_ALIGN.CENTER)

# Set metadata
prs.core_properties.title = "Unified Robot Control Interface"
prs.core_properties.subject = "Agentic Robots, Isaac ROS, CaP-X GUI reference, and Picker2 RL integration"
prs.core_properties.author = "Generated from local project repositories"
prs.core_properties.keywords = "robotics, GUI, Isaac ROS, agentic systems, Picker2, reinforcement learning"
OUT.mkdir(parents=True, exist_ok=True)
prs.save(PPTX)
print(PPTX)
