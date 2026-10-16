#!/usr/bin/env python3
"""Append the Picker2 sim-to-real RL roadmap to the unified robot presentation."""
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
PPT = HERE / "unified_robot_control_interface.pptx"
MEDIA = HERE / "media"
PARALLEL = Path("/home/user/Documents/Code/picker2_rl_grasping/artifacts/parallel_rl_8x8.png")
POSTER = MEDIA / "picker2_rl_simulation_poster.png"
VIDEO = MEDIA / "picker2_rl_simulation.mp4"

BG = RGBColor(0xEF, 0xF7, 0xFF)
NAVY = RGBColor(0x13, 0x2B, 0x44)
MUTED = RGBColor(0x4F, 0x69, 0x82)
TEAL = RGBColor(0x00, 0x89, 0x91)
BLUE = RGBColor(0x23, 0x69, 0xBE)
PURPLE = RGBColor(0x70, 0x43, 0xAF)
GREEN = RGBColor(0x19, 0x87, 0x4F)
GOLD = RGBColor(0xB5, 0x77, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALE = RGBColor(0xDF, 0xEC, 0xF7)


def textbox(slide, x, y, w, h, text, size=16, color=NAVY, bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear(); tf.word_wrap = True; tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]; p.alignment = align
    run = p.add_run(); run.text = text
    run.font.name = "Aptos"; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = color
    return shape


def box(slide, x, y, w, h, fill=WHITE, line=PALE, radius=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    return shape


def header(slide, section, title, number, subtitle=""):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = BG
    textbox(slide, .55, .28, 4.2, .25, section.upper(), 10, TEAL, True)
    textbox(slide, .55, .60, 11.5, .55, title, 27, NAVY, True)
    textbox(slide, 12.25, .34, .5, .25, f"{number:02d}", 9, MUTED, False, PP_ALIGN.RIGHT)
    if subtitle: textbox(slide, .57, 1.16, 11.8, .35, subtitle, 10, MUTED)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.55), Inches(1.50), Inches(12.2), Inches(.025))
    line.fill.solid(); line.fill.fore_color.rgb = PALE; line.line.fill.background()


def add_bullets(slide, x, y, w, h, lines, size=14, color=MUTED):
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(.08); tf.margin_right = Inches(.04)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line; p.level = 0; p.font.name = "Aptos"; p.font.size = Pt(size); p.font.color.rgb = color
        p.space_after = Pt(7); p.text = "• " + p.text
    return sh


def label(slide, x, y, w, text, color):
    sh = box(slide, x, y, w, .32, fill=color, line=color)
    tf=sh.text_frame; tf.clear(); tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
    r=p.add_run(); r.text=text; r.font.name="Aptos"; r.font.size=Pt(9); r.font.bold=True; r.font.color.rgb=WHITE


prs = Presentation(PPT)
if any(
    any(hasattr(shape, "text") and "PICKER2 LEARNING ROADMAP" in shape.text for shape in slide.shapes)
    for slide in prs.slides
):
    raise SystemExit("Picker2 roadmap slides already exist; presentation left unchanged")
blank = prs.slide_layouts[6]

# 05 — roadmap
s = prs.slides.add_slide(blank); header(s, "Picker2 learning roadmap", "From Genesis simulation to the cargo robot", 5,
    "Four validation gates; each stage preserves physical contact and measurable task success")
steps = [
    ("1", "SIMULATION SETUP", TEAL, "Robot + D455 + hook fingertips\nTwo 4 kg suitcase proxies\nCargo platform and 8×8 scene"),
    ("2", "STATE-BASED RL", BLUE, "Privileged geometry + proprioception\nReach → grasp → lift → rotate\nPlace → release → retreat"),
    ("3", "STATE + RGB-D", PURPLE, "Wrist D455 RGB and depth\nVisual encoder + robot state\nDistill from state-based teacher"),
    ("4", "REAL ROBOT", GREEN, "Calibration and safety limits\nDomain randomization\nSlow staged deployment"),
]
for i,(n,t,c,body) in enumerate(steps):
    x=.55+i*3.08; box(s,x,1.88,2.75,3.95); label(s,x+.18,2.08,.52,n,c)
    textbox(s,x+.18,2.58,2.36,.58,t,16,NAVY,True)
    textbox(s,x+.18,3.28,2.36,1.55,body,13,MUTED)
    textbox(s,x+.18,5.18,2.36,.35,"GATE: measurable pass",10,c,True)
    if i<3: textbox(s,x+2.78,3.55,.28,.4,"→",19,MUTED,True,PP_ALIGN.CENTER)
textbox(s,.62,6.34,12,.42,"Final objective: start at home, physically grasp the handle, load flat with handle toward +X, release, and return safe.",14,NAVY,True,PP_ALIGN.CENTER)

# 06 — setup
s = prs.slides.add_slide(blank); header(s,"1 · Simulation setup","One Genesis scene, realistic Picker2 contacts",6,
    "Assets are generated from the Picker2 ROS description; luggage visuals are aligned to collision geometry")
s.shapes.add_picture(str(POSTER), Inches(.55), Inches(1.78), width=Inches(6.15), height=Inches(4.62))
s.shapes.add_picture(str(PARALLEL), Inches(6.92), Inches(1.78), width=Inches(5.85), height=Inches(4.38))
label(s,.72,1.95,1.72,"SINGLE ENVIRONMENT",TEAL); label(s,7.10,1.95,1.55,"8 × 8 TRAINING",BLUE)
textbox(s,.72,5.72,5.72,.53,"Stationary Picker2 base • D455 at the wrist • customized 2FG14 retaining fingertips",11,NAVY,True)
textbox(s,7.10,5.72,5.48,.55,"64 independently resettable environments • 2.2 m spacing • one GPU scene",11,NAVY,True)
textbox(s,.63,6.53,12,.3,"Physics: 4 kg luggage · primitive contact proxies · no fixed grasp constraint · bilateral fingertip contact",12,MUTED,False,PP_ALIGN.CENTER)

# 07 — state RL details
s = prs.slides.add_slide(blank); header(s,"2 · State-based RL","Actions, observations, rewards, and curriculum",7,
    "PPO first learns with privileged state; success is determined by physical luggage state—not trajectory completion")
cols=[(.55,"ACTION · 7D",TEAL,["Δx, Δy, Δz","Δroll, Δpitch, Δyaw","normalized gripper command","bounded at every control step"]),
      (4.34,"OBSERVATION · 63D",BLUE,["46D robot/luggage state","phase target Δposition + quaternion","9D task-phase one-hot + age","TCP, handle, target and velocities"]),
      (8.13,"REWARD + TERMINATION",PURPLE,["target-distance progress","bilateral contact and retention","lift, flatness, target pose, +X handle","release, retreat, success / timeout"])]
for x,title,c,items in cols:
    box(s,x,1.80,3.52,3.72); label(s,x+.18,2.02,2.15,title,c); add_bullets(s,x+.20,2.62,3.08,2.42,items,13)
box(s,.55,5.78,11.10,.78,fill=RGBColor(0xFF,0xF7,0xE5),line=RGBColor(0xE8,0xC9,0x88))
textbox(s,.78,5.96,10.66,.42,"Current engineering status: residual-RL bootstrap reaches grasp/lift/transport; rotation and pure end-to-end RL remain open milestones.",12,GOLD,True)
textbox(s,11.86,5.86,.92,.62,"PPO\nRSL-RL",11,BLUE,True,PP_ALIGN.CENTER)

# 08 — vision policy
s = prs.slides.add_slide(blank); header(s,"3 · Robot state + image","Transition from privileged state to wrist RGB-D",8,
    "Keep proprioception; replace privileged object pose with learned visual features")
for x,title,c,body in [
    (.65,"D455 INPUT",PURPLE,"RGB frame\nDepth / point cloud\nCamera intrinsics\nWrist extrinsics"),
    (3.65,"VISUAL ENCODER",BLUE,"Compact CNN / ViT\nDepth-aware fusion\nHandle + luggage geometry\nTemporal features"),
    (6.65,"POLICY FUSION",TEAL,"Visual latent\nJoint + gripper state\nTCP velocity\nPrevious action"),
    (9.65,"7D CONTROL",GREEN,"Cartesian delta\nOrientation delta\nGripper command\nPhase / memory"),
]:
    box(s,x,2.05,2.35,3.52); label(s,x+.18,2.28,1.72,title,c); textbox(s,x+.22,2.96,1.92,1.78,body,14,MUTED)
for x in (3.08,6.08,9.08): textbox(s,x,3.39,.4,.4,"→",22,NAVY,True,PP_ALIGN.CENTER)
textbox(s,.72,6.02,11.85,.62,"Training recipe: state-policy teacher → synchronized RGB-D rollouts → behavior distillation → PPO fine-tuning with visual randomization.",14,NAVY,True,PP_ALIGN.CENTER)

# 09 — real robot
s = prs.slides.add_slide(blank); header(s,"4 · Transfer to real robot","Sim-to-real deployment with explicit safety gates",9,
    "No direct full-speed deployment: validate sensing, controls, contact, and recovery separately")
items=[
    ("A","CALIBRATE",TEAL,"TCP, D455 extrinsics, platform frame, gripper zero and force limits"),
    ("B","RANDOMIZE",BLUE,"mass, friction, handle pose, camera noise, lighting, latency and controller gains"),
    ("C","SHADOW TEST",PURPLE,"run perception and policy without actuation; compare predicted and measured state"),
    ("D","STAGED MOTION",GOLD,"reach-only → soft grasp → low lift → protected platform placement"),
    ("E","DEPLOY",GREEN,"watchdog, workspace clamps, contact abort, human E-stop and automatic retreat"),
]
for i,(n,t,c,b) in enumerate(items):
    y=1.78+i*.96; label(s,.66,y,.44,n,c); textbox(s,1.28,y-.02,2.05,.34,t,14,NAVY,True); textbox(s,3.40,y-.02,8.74,.58,b,13,MUTED)
textbox(s,.72,6.66,11.8,.27,"Release criterion: repeated randomized simulation success + hardware reach/grasp validation + bounded failure behavior.",11,NAVY,True,PP_ALIGN.CENTER)

# 10 — embedded demo
s = prs.slides.add_slide(blank); header(s,"Simulation evidence","Current Picker2 policy playback",10,
    "Embedded 35-second recording from the local Genesis environment (click Play in presentation mode)")
s.shapes.add_movie(str(VIDEO), Inches(.72), Inches(1.78), Inches(8.52), Inches(4.79),
                   poster_frame_image=str(POSTER), mime_type="video/mp4")
box(s,9.55,1.78,3.05,4.79)
label(s,9.78,2.04,1.52,"WHAT TO WATCH",TEAL)
add_bullets(s,9.78,2.65,2.54,2.65,["home-position start","handle approach","physical grasp","retained lift","transport toward platform","current rotation limitation"],12)
textbox(s,9.78,5.84,2.52,.47,"Video is evidence of the current hybrid residual-RL stage—not a claim of solved end-to-end RL.",9,GOLD,True)

backup = PPT.with_name(PPT.stem + ".before_picker2_rl.pptx")
if not backup.exists(): backup.write_bytes(PPT.read_bytes())
prs.save(PPT)
print(f"updated {PPT} ({len(prs.slides)} slides)")
